# -*- coding: utf-8 -*-
"""Detecção de gênero/número, construção de contexto e flexão de texto."""

import docx

import organizar_processos as op
from organizar_processos import Contexto, Parte, Processo


# --------------------------------------------------------------------------
# _strip_accents / _detectar_genero
# --------------------------------------------------------------------------


def test_strip_accents():
    assert op._strip_accents("Carapicuíba") == "Carapicuiba"
    assert op._strip_accents("São Paulo") == "Sao Paulo"


def test_detectar_genero_feminino_por_terminacao():
    assert op._detectar_genero("Maria Silva", "Pessoa Física") == "f"


def test_detectar_genero_masculino_por_terminacao():
    assert op._detectar_genero("João Souza", "Pessoa Física") == "m"


def test_detectar_genero_override_masculino():
    # "Lucas" termina em 's' mas há override; e nomes em 'a' com override.
    assert op._detectar_genero("Lucas Lima", "Pessoa Física") == "m"


def test_detectar_genero_override_feminino():
    assert op._detectar_genero("Beatriz Costa", "Pessoa Física") == "f"


def test_detectar_genero_pj_sempre_masculino():
    assert op._detectar_genero("Empresa Qualquer", "Pessoa Jurídica") == "m"


def test_detectar_genero_vazio():
    assert op._detectar_genero("", "") == "m"


# --------------------------------------------------------------------------
# _construir_contexto
# --------------------------------------------------------------------------


def test_construir_contexto_singular():
    proc = Processo(partes=[
        Parte(papel="AUTOR", nome="Maria Silva", tipo="Pessoa Física"),
        Parte(papel="RÉU", nome="João Souza", tipo="Pessoa Física"),
    ])
    ctx = op._construir_contexto(proc)
    assert ctx.genero_autor == "f"
    assert ctx.numero_autor == "s"
    assert ctx.genero_reu == "m"
    assert ctx.numero_reu == "s"


def test_construir_contexto_plural_misto_vence_masculino():
    proc = Processo(partes=[
        Parte(papel="RÉU", nome="Ana Lima", tipo="Pessoa Física"),
        Parte(papel="RÉU", nome="João Souza", tipo="Pessoa Física"),
    ])
    ctx = op._construir_contexto(proc)
    # Plural com pelo menos um masculino -> masculino plural.
    assert ctx.genero_reu == "m"
    assert ctx.numero_reu == "p"


def test_construir_contexto_sem_partes_default_masculino_singular():
    ctx = op._construir_contexto(Processo())
    assert (ctx.genero_autor, ctx.numero_autor) == ("m", "s")
    assert (ctx.genero_reu, ctx.numero_reu) == ("m", "s")


# --------------------------------------------------------------------------
# _inflexionar_texto
# --------------------------------------------------------------------------


def test_inflexionar_forma_isolada():
    ctx = Contexto(genero_autor="f", numero_autor="s")
    assert op._inflexionar_texto("intimar o autor", ctx) == "intimar a autora"


def test_inflexionar_prep_e_forma():
    ctx = Contexto(genero_reu="m", numero_reu="p")
    # "o executado" -> "os executados"
    assert op._inflexionar_texto("citar o executado", ctx) == "citar os executados"


def test_inflexionar_preserva_caixa():
    ctx = Contexto(genero_autor="f", numero_autor="s")
    assert op._inflexionar_texto("O Autor requer", ctx) == "A Autora requer"


def test_inflexionar_pula_citacao_legal():
    ctx = Contexto(genero_autor="f", numero_autor="s")
    texto = "conforme art. 5º o autor"
    # Parágrafo com citação legal é preservado integralmente.
    assert op._inflexionar_texto(texto, ctx) == texto


def test_inflexionar_sem_contexto_retorna_original():
    assert op._inflexionar_texto("o autor", None) == "o autor"


# --------------------------------------------------------------------------
# _detectar_genero_juiz
# --------------------------------------------------------------------------


def test_detectar_genero_juiz_feminino():
    assert op._detectar_genero_juiz("Maria Silva") == "f"
    assert op._detectar_genero_juiz("Ana Souza") == "f"


def test_detectar_genero_juiz_masculino():
    assert op._detectar_genero_juiz("João Souza") == "m"
    assert op._detectar_genero_juiz("Carlos Silva") == "m"


def test_detectar_genero_juiz_override():
    assert op._detectar_genero_juiz("Beatriz Costa") == "f"  # termina em 'z' mas override feminino
    assert op._detectar_genero_juiz("Lucas Lima") == "m"     # termina em 's' mas override masculino


def test_detectar_genero_juiz_vazio():
    assert op._detectar_genero_juiz("") == "m"


# --------------------------------------------------------------------------
# _construir_contexto com juiz
# --------------------------------------------------------------------------


def test_construir_contexto_com_juiz():
    proc = Processo(
        partes=[
            Parte(papel="AUTOR", nome="Maria Silva", tipo="Pessoa Física"),
            Parte(papel="RÉU", nome="João Souza", tipo="Pessoa Física"),
        ],
        juiz="Maria Silva"
    )
    ctx = op._construir_contexto(proc)
    assert ctx.genero_autor == "f"
    assert ctx.numero_autor == "s"
    assert ctx.genero_reu == "m"
    assert ctx.numero_reu == "s"
    assert ctx.genero_juiz == "f"  # Juiz feminino


def test_construir_contexto_com_juiz_masculino():
    proc = Processo(
        partes=[
            Parte(papel="AUTOR", nome="Maria Silva", tipo="Pessoa Física"),
            Parte(papel="RÉU", nome="João Souza", tipo="Pessoa Física"),
        ],
        juiz="João Souza"
    )
    ctx = op._construir_contexto(proc)
    assert ctx.genero_juiz == "m"  # Juiz masculino


def test_construir_contexto_sem_juiz_default_masculino():
    proc = Processo(
        partes=[
            Parte(papel="AUTOR", nome="Maria Silva", tipo="Pessoa Física"),
            Parte(papel="RÉU", nome="João Souza", tipo="Pessoa Física"),
        ],
        juiz=""  # Juiz vazio
    )
    ctx = op._construir_contexto(proc)
    assert ctx.genero_juiz == "m"  # Default para masculino


# --------------------------------------------------------------------------
# _flex_judicial_treatment
# --------------------------------------------------------------------------


def test_flex_judicial_treatment_excelentissimo_senhor_doutor_juiz_masculino():
    ctx = Contexto(genero_juiz="m")
    texto = "EXCELENTÍSSIMO SENHOR DOUTOR JUIZ"
    resultado = op._flex_judicial_treatment(texto, ctx)
    assert resultado == "EXCELENTÍSSIMO SENHOR DOUTOR JUIZ"  # Masculino permanece igual


def test_flex_judicial_treatment_excelentissimo_senhora_doutora_juiza_feminino():
    ctx = Contexto(genero_juiz="f")
    texto = "EXCELENTÍSSIMO SENHOR DOUTOR JUIZ"
    resultado = op._flex_judicial_treatment(texto, ctx)
    assert resultado == "EXCELENTÍSSIMA SENHORA DOUTORA JUIZA"  # Feminino adiciona 'A'


def test_flex_judicial_treatment_meritissimo_senhor_doutor_juiz_masculino():
    ctx = Contexto(genero_juiz="m")
    texto = "MERITÍSSIMO SENHOR DOUTOR JUIZ"
    resultado = op._flex_judicial_treatment(texto, ctx)
    assert resultado == "MERITÍSSIMO SENHOR DOUTOR JUIZ"  # Masculino permanece igual


def test_flex_judicial_treatment_meritissima_senhora_doutora_juiza_feminino():
    ctx = Contexto(genero_juiz="f")
    texto = "MERITÍSSIMO SENHOR DOUTOR JUIZ"
    resultado = op._flex_judicial_treatment(texto, ctx)
    assert resultado == "MERITÍSSIMA SENHORA DOUTORA JUIZA"  # Feminino adiciona 'A'


def test_flex_judicial_treatment_com_espacos_multiplos():
    ctx = Contexto(genero_juiz="f")
    texto = "EXCELENTÍSSIMO   SENHOR      DOUTOR    JUIZ"
    resultado = op._flex_judicial_treatment(texto, ctx)
    assert resultado == "EXCELENTÍSSIMA SENHORA DOUTORA JUIZA"  # Normaliza espaços e adiciona 'A'


def test_flex_judicial_treatment_sem_contexto():
    texto = "EXCELENTÍSSIMO SENHOR DOUTOR JUIZ"
    resultado = op._flex_judicial_treatment(texto, None)
    assert resultado == "EXCELENTÍSSIMO SENHOR DOUTOR JUIZ"  # Deve permanecer inalterado


def test_flex_judicial_treatment_nao_correponde():
    ctx = Contexto(genero_juiz="f")
    texto = "EXCELENTÍSSIMO DR. JOAO SILVA"
    resultado = op._flex_judicial_treatment(texto, ctx)
    assert resultado == "EXCELENTÍSSIMO DR. JOAO SILVA"  # Não deve ser alterado


# --------------------------------------------------------------------------
# Tratamento judicial com parênteses (A)
# --------------------------------------------------------------------------


def test_flex_judicial_treatment_parenteses_masculino():
    ctx = Contexto(genero_juiz="m")
    texto = "EXCELENTÍSSIMO(A) SENHOR(A) DOUTOR(A) JUIZ(A)"
    resultado = op._flex_judicial_treatment(texto, ctx)
    assert resultado == "EXCELENTÍSSIMO SENHOR DOUTOR JUIZ"


def test_flex_judicial_treatment_parenteses_feminino():
    ctx = Contexto(genero_juiz="f")
    texto = "EXCELENTÍSSIMO(A) SENHOR(A) DOUTOR(A) JUIZ(A)"
    resultado = op._flex_judicial_treatment(texto, ctx)
    assert resultado == "EXCELENTÍSSIMA SENHORA DOUTORA JUIZA"


def test_flex_judicial_treatment_parenteses_meritissimo_masculino():
    ctx = Contexto(genero_juiz="m")
    texto = "MERITÍSSIMO(A) SENHOR(A) DOUTOR(A) JUIZ(A)"
    resultado = op._flex_judicial_treatment(texto, ctx)
    assert resultado == "MERITÍSSIMO SENHOR DOUTOR JUIZ"


def test_flex_judicial_treatment_parenteses_meritissimo_feminino():
    ctx = Contexto(genero_juiz="f")
    texto = "MERITÍSSIMO(A) SENHOR(A) DOUTOR(A) JUIZ(A)"
    resultado = op._flex_judicial_treatment(texto, ctx)
    assert resultado == "MERITÍSSIMA SENHORA DOUTORA JUIZA"


def test_flex_judicial_treatment_parenteses_misto():
    ctx = Contexto(genero_juiz="f")
    texto = "EXCELENTÍSSIMO(A) SENHOR(A) DOUTOR JUIZ"
    resultado = op._flex_judicial_treatment(texto, ctx)
    assert resultado == "EXCELENTÍSSIMA SENHORA DOUTORA JUIZA"


def test_flex_judicial_treatment_sem_doutor_masculino():
    ctx = Contexto(genero_juiz="m")
    texto = "EXCELENTÍSSIMO(A) SENHOR(A) JUIZ(A)"
    resultado = op._flex_judicial_treatment(texto, ctx)
    assert resultado == "EXCELENTÍSSIMO SENHOR JUIZ"


def test_flex_judicial_treatment_sem_doutor_feminino():
    ctx = Contexto(genero_juiz="f")
    texto = "EXCELENTÍSSIMO(A) SENHOR(A) JUIZ(A)"
    resultado = op._flex_judicial_treatment(texto, ctx)
    assert resultado == "EXCELENTÍSSIMA SENHORA JUIZA"


def test_flex_judicial_treatment_sem_doutor_meritissimo_feminino():
    ctx = Contexto(genero_juiz="f")
    texto = "MERITÍSSIMO(A) SENHOR(A) JUIZ(A)"
    resultado = op._flex_judicial_treatment(texto, ctx)
    assert resultado == "MERITÍSSIMA SENHORA JUIZA"


def test_inflexionar_texto_com_tratamento_judicial_parenteses_feminino():
    ctx = Contexto(genero_juiz="f", genero_autor="f", numero_autor="s")
    texto = "EXCELENTÍSSIMO(A) SENHOR(A) DOUTOR(A) JUIZ(A). VEM REQUERER O AUTOR."
    resultado = op._inflexionar_texto(texto, ctx)
    assert resultado == "EXCELENTÍSSIMA SENHORA DOUTORA JUIZA. VEM REQUERER A AUTORA."


def test_inflexionar_texto_com_tratamento_judicial_parenteses_masculino():
    ctx = Contexto(genero_juiz="m", genero_autor="f", numero_autor="s")
    texto = "EXCELENTÍSSIMO(A) SENHOR(A) DOUTOR(A) JUIZ(A). VEM REQUERER O AUTOR."
    resultado = op._inflexionar_texto(texto, ctx)
    assert resultado == "EXCELENTÍSSIMO SENHOR DOUTOR JUIZ. VEM REQUERER A AUTORA."


# --------------------------------------------------------------------------
# _inflexionar_texto com tratamento judicial
# --------------------------------------------------------------------------


def test_inflexionar_texto_com_tratamento_judicial_masculino():
    ctx = Contexto(genero_juiz="m", genero_autor="f", numero_autor="s")
    texto = "EXCELENTÍSSIMO SENHOR DOUTOR JUIZ. VEM REQUERER O AUTOR."
    resultado = op._inflexionar_texto(texto, ctx)
    assert resultado == "EXCELENTÍSSIMO SENHOR DOUTOR JUIZ. VEM REQUERER A AUTORA."


def test_inflexionar_texto_com_tratamento_judicial_feminino():
    ctx = Contexto(genero_juiz="f", genero_autor="f", numero_autor="s")
    texto = "EXCELENTÍSSIMO SENHOR DOUTOR JUIZ. VEM REQUERER O AUTOR."
    resultado = op._inflexionar_texto(texto, ctx)
    assert resultado == "EXCELENTÍSSIMA SENHORA DOUTORA JUIZA. VEM REQUERER A AUTORA."


def test_inflexionar_texto_tratamento_judicial_no_closing():
    ctx = Contexto(genero_juiz="f", genero_autor="m", numero_autor="s")
    texto = "EXCELENTÍSSIMO SENHOR DOUTOR JUIZ\n\nVEM REQUERER O AUTOR\n\nTERMOS EM QUE PEDE DEFERIMENTO."
    resultado = op._inflexionar_texto(texto, ctx)
    assert "EXCELENTÍSSIMA SENHORA DOUTORA JUIZA" in resultado
    assert "VEM REQUERER O AUTOR" in resultado  # Autor masculino permanece
    assert "TERMOS EM QUE PEDE DEFERIMENTO" in resultado


# --------------------------------------------------------------------------
# _classificar_paragrafos / _norm_paragrafo
# --------------------------------------------------------------------------


def test_norm_paragrafo():
    assert op._norm_paragrafo("  Texto  Exemplo.  ") == "texto exemplo"


def test_classificar_paragrafos_header_body_closing():
    doc = docx.Document()
    doc.add_paragraph("EXCELENTÍSSIMO SENHOR DOUTOR JUIZ")
    doc.add_paragraph("")
    doc.add_paragraph("Vem requerer o que segue.")
    doc.add_paragraph("Termos em que pede deferimento.")
    zonas = op._classificar_paragrafos(doc)
    assert zonas[0] == "HEADER"
    assert zonas[2] == "BODY"
    assert zonas[3] == "CLOSING"
