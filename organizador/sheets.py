# -*- coding: utf-8 -*-
"""Leitura das planilhas Google do escritório (somente leitura).

Duas coisas moram aqui, e as duas usam o mesmo acesso:

* os PROCESSOS das planilhas de campanha (`coletar_processos_das_planilhas`);
* as SENHAS dos tribunais (`coletar_credenciais_da_planilha`), fonte do botão
  "Atualizar senhas dos tribunais" — senha trocada no escritório chega em todo
  mundo com um clique, sem release novo.

Adaptado do ``sync_sheets.py`` de referência: usa uma CONTA DE SERVIÇO
(``token_google_api.json`` na raiz, apontado por ``GOOGLE_SA_PATH``) com escopos
*.readonly — as planilhas NUNCA são alteradas. Cada planilha precisa estar
compartilhada como *Leitor* com o e-mail da conta de serviço (impresso ao rodar).

Suporta planilhas nativas do Google (exporta .xlsx) e .xlsx subidos ao Drive,
com múltiplas abas (escolhe a de mais CNJs) e cabeçalho deslocado. Devolve os
números já no formato canônico do app (``RE_NUMERO_CNJ``), deduplicados.

Uso programático:
    from organizador import sheets
    regs = sheets.coletar_processos_das_planilhas()   # [{"numero", "campanha"}]
"""
from __future__ import annotations

import io
import os
import re
import unicodedata
import warnings
from typing import Callable

from .config import GOOGLE_SA_PATH, _login_credencial, _valor_credencial
from .padroes import RE_NUMERO_CNJ

SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets.readonly",
    "https://www.googleapis.com/auth/drive.readonly",
]
MIME_GSHEET = "application/vnd.google-apps.spreadsheet"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_HDR = 8  # linhas do topo em que procuramos o cabeçalho (cabeçalho deslocado)
MAX_LINHAS = 50_000  # teto por aba (protege a RAM em máquinas de 4 GB)

# campanha (rótulo) -> (file_id do Drive, dica do nome da coluna de processo).
# A dica é só um chute para achar a linha de cabeçalho; a detecção da coluna é
# tolerante (contém 'processo'/'cnj' e, no limite, a coluna com mais CNJs), então
# pequenas variações no cabeçalho não quebram a leitura.
SHEETS: dict[str, tuple[str, str]] = {
    "EMCCAMP Judicial":           ("1GZ1N-dBF1G9FcK_9qj7t-MBy-eEhZW25", "Nº PROCESSO"),
    "Longitude Judicial":         ("1E-WgMiPClhPq9hb6PXvo7PqkCHmZgnq6f2pk2RoEkRw", "Nº PROCESSO"),
    "Cyrela Judicial":            ("17sHrE8igYgTgXHmQVlnbLXKbod6RvOzPq5CfqbMOPC4", "PROCESSO"),
    "Arbore - Judicial":          ("1VbYKZo4tM9_Tg0R68vwpRi4FO2_xVJi_", "Nº PROCESSO"),
    "RSF Judicial":               ("1y3Mm1xkOPLR2BT2uUvciKm6ms0ItB5kY", "Nº PROCESSO"),
    "Movida Judicial":            ("1ECNr7qtztpZPLaXWogGazk9PeO7NhknYS3Q3nsj0MF8", "Nº PROCESSO"),
    "Grupo Super Nosso Judicial": ("1PkdyMHHAXe8w3yeFhmXZOF9kLckoLcRvJk9R6Asm3EA", "PROCESSO"),
    "VIC Judicial":               ("1wlpxK0RZGkLevObMwkPhe-vQQUnco7NwEBBkcvjVct4", "Processo"),
    "BVEP":                       ("1c4_qfDnS0b2O7kwgTz5BJI4ZU10GQJgi", "Numero Processual"),
    "São José Judicial":          ("1n6HDw79nF_m8p0xAI0lWrCkVecRTzZJI-2GXGJr8M_A", "PROCESSO"),
    "Level/Cantupneus":           ("1X8WH_u75OuqrczNlgbFJRdZbAYY99qc5RMBN8LYTIYA", "PROCESSO"),
}

# --- planilha de senhas dos tribunais --------------------------------------
MIN_LINHAS_SENHAS = 3  # abaixo disso é capa/dashboard, não a tabela de senhas
# Sinônimos em ordem MAIS-ESPECÍFICO-PRIMEIRO: o casamento por "contém" faria
# "senha" ganhar de "senha nova" e levar a coluna errada (ex.: "Senha antiga").
SIN_TJ = ("tribunal/sistema", "tj/sistema", "tribunal", "sistema", "tj",
          "orgao", "portal", "site", "descricao")
SIN_LOGIN = ("login", "usuario", "user", "cpf/cnpj", "cpf", "e-mail", "email")
SIN_SENHA = ("senha nova", "senha atual", "senha", "password", "pass")


def _id_de_url(valor: str) -> str:
    """Aceita o id cru OU a URL completa da planilha."""
    m = re.search(r"/spreadsheets/d/([A-Za-z0-9_-]+)", str(valor or ""))
    return m.group(1) if m else str(valor or "").strip()


# Escape hatch por variável de ambiente (aceita URL colada): se a planilha for
# recriada o id muda, e sem esta saída a correção exigiria um .exe novo — que é
# exatamente o problema que ler a planilha veio matar.
SHEET_SENHAS: str = _id_de_url(
    os.getenv("ORGANIZADOR_SENHAS_SHEET_ID")
    or "1BkNU0gkcl9A3O88i10e1Qbih4Abf36nthfpIj478bZg"
)


# ---------------------------------------------------------------------
# Acesso ao Google (somente leitura) — import tardio do google-auth
# ---------------------------------------------------------------------
def _sessao():
    if not GOOGLE_SA_PATH.exists():
        raise RuntimeError(
            f"credencial da conta de serviço não encontrada em {GOOGLE_SA_PATH} "
            "(defina GOOGLE_CREDENTIALS_FILE ou coloque token_google_api.json na raiz)."
        )
    from google.oauth2.service_account import Credentials
    from google.auth.transport.requests import AuthorizedSession

    creds = Credentials.from_service_account_file(str(GOOGLE_SA_PATH), scopes=SCOPES)
    return AuthorizedSession(creds), creds.service_account_email


# ---------------------------------------------------------------------
# Normalização de nomes/valores
# ---------------------------------------------------------------------
def _norm(s) -> str:
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().casefold()


def _normalizar_cnj(texto) -> str:
    """Devolve o CNJ no formato canônico do app (RE_NUMERO_CNJ) ou '' se a célula
    não contém um número de processo. Aceita o canônico pontuado e também 20
    dígitos crus (reconstrói a pontuação)."""
    s = str(texto or "")
    m = RE_NUMERO_CNJ.search(s)
    if m:
        return m.group(0)
    d = re.sub(r"\D", "", s)
    if len(d) == 20:
        return f"{d[:7]}-{d[7:9]}.{d[9:13]}.{d[13]}.{d[14:16]}.{d[16:20]}"
    return ""


def _achar_col(
    cabecalho: list[str], alvo: str, excluir: frozenset[int] | set[int] = frozenset()
) -> int | None:
    """Índice da coluna cujo rótulo casa `alvo` (igual; senão, contém).

    `excluir` tira índices da disputa — é o que impede, na leitura de senhas, que
    tj/login/senha caiam todos na mesma coluna quando os sinônimos se sobrepõem.
    """
    alvo_n = _norm(alvo)
    normalizados = [_norm(c) for c in cabecalho]
    for i, c in enumerate(normalizados):
        if i not in excluir and c == alvo_n:
            return i
    for i, c in enumerate(normalizados):  # fallback: contém
        # Cabeçalho VAZIO casava qualquer alvo (`"" in alvo_n` é True) e elegia
        # uma coluna em branco — comum em planilha exportada. Só não explodia
        # porque `_detectar_col_processo` regatava o acerto com `conta_cnj > 0`;
        # a leitura de senhas não tem esse contador para se defender.
        if i in excluir or not c:
            continue
        if alvo_n and (alvo_n in c or c in alvo_n):
            return i
    return None


def _achar_col_sinonimos(
    cabecalho: list[str],
    sinonimos: tuple[str, ...],
    excluir: frozenset[int] | set[int] = frozenset(),
) -> int | None:
    """Primeiro sinônimo que casa uma coluna. A ORDEM da tupla importa: o mais
    específico vem antes (`senha nova` antes de `senha`), senão o genérico casa
    primeiro por 'contém' e leva a coluna errada."""
    for alvo in sinonimos:
        i = _achar_col(cabecalho, alvo, excluir)
        if i is not None:
            return i
    return None


def _detectar_col_processo(
    cabecalho: list[str], conta_cnj: Callable[[int], int], hint: str
) -> int | None:
    """Coluna de número de processo: 1) casa a dica; 2) cabeçalho contém
    'processo'/'cnj'; 3) a coluna com mais valores em formato CNJ.

    `conta_cnj(i)` devolve quantos CNJs há na coluna `i` ABAIXO do cabeçalho."""
    i = _achar_col(cabecalho, hint)
    if i is not None and conta_cnj(i) > 0:
        return i
    for i, rotulo in enumerate(cabecalho):
        cab = _norm(rotulo)
        if ("processo" in cab or "cnj" in cab) and conta_cnj(i) > 0:
            return i
    melhor, melhor_q = None, 0
    for i in range(len(cabecalho)):
        q = conta_cnj(i)
        if q > melhor_q:
            melhor, melhor_q = i, q
    return melhor if melhor_q >= 2 else None


# ---------------------------------------------------------------------
# Download + leitura da planilha
# ---------------------------------------------------------------------
def _baixar_bytes(s, fid: str) -> io.BytesIO:
    """Baixa o binário .xlsx (planilha nativa é exportada; arquivo é get_media)."""
    meta = s.get(
        f"https://www.googleapis.com/drive/v3/files/{fid}",
        params={"fields": "name,mimeType", "supportsAllDrives": "true"},
        timeout=30,
    ).json()
    if "error" in meta:
        raise RuntimeError(meta["error"].get("message", str(meta["error"]))[:120])
    if meta.get("mimeType") == MIME_GSHEET:
        r = s.get(
            f"https://www.googleapis.com/drive/v3/files/{fid}/export",
            params={"mimeType": MIME_XLSX},
            timeout=30,
        )
    else:
        r = s.get(
            f"https://www.googleapis.com/drive/v3/files/{fid}",
            params={"alt": "media", "supportsAllDrives": "true"},
            timeout=30,
        )
    r.raise_for_status()
    return io.BytesIO(r.content)


def _valor_celula(v) -> str:
    """Célula → str, sem notação científica para inteiros exatos.

    CUIDADO com o limite do float: o openpyxl devolve célula numérica de 20
    dígitos como ``float``, e um float64 só representa inteiros exatos até
    2**53. Converter ``4.008846102025826e+19`` com ``int()`` produziria um número
    plausível mas ERRADO, que o ``_normalizar_cnj`` aceitaria como CNJ válido —
    pior que não ler nada. Acima do limite devolvemos a forma científica, que
    não casa como CNJ (é o que o pandas fazia, e por isso a troca de motor não
    mudou nenhum resultado nas 11 planilhas reais).

    Na prática os CNJs das planilhas do escritório são TEXTO; este ramo é defesa.
    """
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) and v.is_integer():
        return str(int(v)) if abs(v) < 2 ** 53 else repr(v)
    return str(v)


def _grade(ws, log: Callable[[str], None]) -> list[list[str]]:
    """Materializa a aba como lista de linhas de str, UMA passada só."""
    grade: list[list[str]] = []
    for i, linha in enumerate(ws.iter_rows(values_only=True)):
        if i >= MAX_LINHAS:
            log(f"   ⚠️ aba '{ws.title}': truncada em {MAX_LINHAS} linhas.")
            break
        grade.append([_valor_celula(v) for v in linha])
    return grade


def _mapa_cnj(grade: list[list[str]]) -> tuple[int, list[int], list[list[int]]]:
    """Pré-computa, por coluna: o total de CNJs e quantos estão nas primeiras
    linhas (prefixo até MAX_HDR).

    Isto é o que substitui o custo do antigo laço com pandas: a regex de CNJ roda
    UMA vez por célula, e a contagem para cada cabeçalho candidato (0..MAX_HDR)
    sai por subtração de prefixo, em vez de reparsear a aba 9 vezes."""
    n_col = max((len(l) for l in grade), default=0)
    total = [0] * n_col
    # prefixo[i][h] = CNJs da coluna i nas primeiras h linhas.
    limite = min(MAX_HDR + 1, len(grade) + 1)
    prefixo = [[0] * limite for _ in range(n_col)]
    for l, linha in enumerate(grade):
        for i, cel in enumerate(linha):
            if not cel or not _normalizar_cnj(cel):
                continue
            total[i] += 1
            for h in range(l + 1, limite):
                prefixo[i][h] += 1
    return n_col, total, prefixo


def ler_planilha(s, fid: str, hint: str, log: Callable[[str], None] = print):
    """Varre TODAS as abas e, em cada uma, TODAS as linhas de cabeçalho plausíveis
    (0..MAX_HDR), escolhendo o par (aba, cabeçalho) cuja coluna de processo tem
    mais CNJs válidos. Devolve ``(valores_da_coluna, rótulo_da_coluna)``.

    A detecção é dirigida por CONTAGEM DE CNJ (não pelo texto do cabeçalho), então
    funciona com cabeçalhos fora do padrão (ex.: BVEP usa 'Numero Processual', que
    não contém a palavra 'processo') e com cabeçalho deslocado/capas/dashboards.

    Lê com openpyxl em modo ``read_only`` (streaming) — sem pandas, que arrastava
    numpy/pyarrow para o bundle e reparseava cada aba até 9 vezes."""
    from openpyxl import load_workbook

    dados = _baixar_bytes(s, fid)
    with warnings.catch_warnings():
        # As planilhas exportadas do Google trazem extensões que o openpyxl não
        # conhece; o aviso é ruído no log da GUI.
        warnings.simplefilter("ignore", UserWarning)
        wb = load_workbook(dados, read_only=True, data_only=True)
        try:
            abas = list(wb.sheetnames)
            melhor: tuple[list[str], str] | None = None
            melhor_n = -1
            for nome in abas:
                try:
                    grade = _grade(wb[nome], log)
                except Exception:
                    continue
                if not grade:
                    continue
                _n_col, total, prefixo = _mapa_cnj(grade)
                for hdr in range(min(MAX_HDR, len(grade))):
                    cabecalho = grade[hdr]

                    def conta_cnj(i: int, _hdr: int = hdr) -> int:
                        if i >= len(total):
                            return 0
                        return total[i] - prefixo[i][_hdr + 1]

                    col = _detectar_col_processo(cabecalho, conta_cnj, hint)
                    if col is None:
                        continue
                    n = conta_cnj(col)
                    if n > melhor_n:
                        rotulo = cabecalho[col] if col < len(cabecalho) else f"col {col}"
                        valores = [
                            linha[col] if col < len(linha) else ""
                            for linha in grade[hdr + 1:]
                        ]
                        melhor, melhor_n = (valores, rotulo), n
        finally:
            wb.close()
    if melhor is None:
        raise RuntimeError(f"nenhuma aba com coluna de processo (abas: {abas})")
    return melhor


# ---------------------------------------------------------------------
# API principal
# ---------------------------------------------------------------------
def coletar_processos_das_planilhas(
    log: Callable[[str], None] = print,
) -> list[dict]:
    """Lê todas as planilhas de SHEETS e devolve ``[{"numero", "campanha"}]``
    deduplicado por CNJ (1ª campanha vence). Somente leitura."""
    s, robot = _sessao()
    log(f"🔑 conta de serviço: {robot} (compartilhe cada planilha como Leitor com ele)")
    registros: list[dict] = []
    vistos: set[str] = set()
    ok = 0
    primeiro_erro = ""
    for campanha, (fid, hint) in SHEETS.items():
        try:
            valores, col = ler_planilha(s, fid, hint, log=log)
        except Exception as exc:
            if not primeiro_erro:
                primeiro_erro = str(exc)
            log(f"   ❌ {campanha}: {str(exc)[:100]}")
            continue
        ok += 1
        n = 0
        for v in valores:
            cnj = _normalizar_cnj(v)
            if not cnj or cnj in vistos:
                continue
            vistos.add(cnj)
            registros.append({"numero": cnj, "campanha": campanha})
            n += 1
        log(f"   📄 {campanha}: {n} processos (coluna '{col}')")
    # Se NENHUMA planilha pôde ser lida, é falha de config (API desabilitada,
    # credencial sem acesso, planilhas não compartilhadas) — não um "em dia".
    # Estoura para o chamador (GUI/CLI) mostrar o erro real em vez de sucesso vazio.
    if ok == 0 and SHEETS:
        raise RuntimeError(
            "nenhuma planilha pôde ser lida — verifique se as APIs Google Drive/"
            f"Sheets estão habilitadas no projeto e se cada planilha foi compartilhada "
            f"como Leitor com {robot}. Primeiro erro: {primeiro_erro[:160]}"
        )
    log(f"📋 total: {len(registros)} processos únicos nas planilhas.")
    return registros


# ---------------------------------------------------------------------
# Senhas dos tribunais
# ---------------------------------------------------------------------
def _linhas_de_senha(
    linhas: list[list[str]], i_tj: int, i_login: int, i_senha: int
) -> list[dict[str, str]]:
    """Linhas da tabela → ``[{"tj", "login", "senha"}]``, na ordem da planilha.

    Linha válida = tribunal preenchido E (login ou senha). Só o tribunal
    preenchido costuma ser título de seção ("TRIBUNAIS ESTADUAIS"), não acesso.

    NÃO tenta separar "lixo" de credencial: o arquivo real carrega linhas de
    recado ("senha robo", "SENHA CERTIFICADO A1: ...") que são informação do
    escritório. Classificá-las faria elas sumirem da máquina de quem sincroniza.
    """
    def cel(linha: list[str], i: int) -> str:
        return linha[i] if i < len(linha) else ""  # linha curta não é IndexError

    itens: list[dict[str, str]] = []
    for linha in linhas:
        tj = _valor_credencial(cel(linha, i_tj))
        if not tj:
            continue
        login = _login_credencial(cel(linha, i_login))
        senha = _valor_credencial(cel(linha, i_senha))
        if not login and not senha:
            continue
        itens.append({"tj": tj, "login": login, "senha": senha})
    return itens


def ler_senhas_da_planilha(
    s, fid: str, log: Callable[[str], None] = print
) -> tuple[list[dict[str, str]], str]:
    """Lê a planilha de senhas. Devolve ``(itens, descrição_da_detecção)``.

    Mesma estratégia de `ler_planilha` — varre TODAS as abas e TODAS as linhas de
    cabeçalho plausíveis (0..MAX_HDR) —, mas resolvendo um TRIO de colunas
    (tribunal, login, senha) em vez de uma. Sem a contagem de CNJ para se guiar,
    a nota do candidato é ``(linhas_válidas, linhas_com_senha)``: a tabela de
    verdade é a coisa mais larga do arquivo, e o desempate por senha preenchida
    impede que uma aba de resumo "Tribunal | Login" ganhe da tabela real.

    Ordem de resolução senha → tribunal → login, cada uma excluindo os índices já
    tomados: os sinônimos se sobrepõem (uma coluna "Descrição" casaria tribunal e
    login), e sem exclusão os três cairiam na mesma coluna.

    A descrição devolvida traz aba, linha de cabeçalho e RÓTULOS das colunas —
    é o que o usuário confere antes de aceitar. Nunca contém valor de senha.
    """
    from openpyxl import load_workbook

    dados = _baixar_bytes(s, fid)
    abas: list[str] = []
    melhor: tuple[list[dict[str, str]], str] | None = None
    melhor_nota = (0, 0)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        wb = load_workbook(dados, read_only=True, data_only=True)
        try:
            abas = list(wb.sheetnames)
            for nome in abas:
                try:
                    grade = _grade(wb[nome], log)
                except Exception:
                    continue
                for hdr in range(min(MAX_HDR, len(grade))):
                    cab = grade[hdr]
                    i_senha = _achar_col_sinonimos(cab, SIN_SENHA)
                    if i_senha is None:
                        continue
                    i_tj = _achar_col_sinonimos(cab, SIN_TJ, {i_senha})
                    if i_tj is None:
                        continue
                    i_login = _achar_col_sinonimos(cab, SIN_LOGIN, {i_senha, i_tj})
                    if i_login is None:
                        continue
                    itens = _linhas_de_senha(grade[hdr + 1:], i_tj, i_login, i_senha)
                    n_senha = sum(1 for i in itens if i["senha"])
                    # Portões: tabela curta é capa/dashboard; coluna de senha
                    # inteiramente vazia é aba de resumo. Aplicar qualquer uma
                    # das duas apagaria senha boa por campo vazio.
                    if len(itens) < MIN_LINHAS_SENHAS or not n_senha:
                        continue
                    nota = (len(itens), n_senha)
                    if nota > melhor_nota:  # `>` estrito: no empate o 1º vence
                        melhor_nota = nota
                        melhor = (
                            itens,
                            f"aba '{nome}', cabeçalho na linha {hdr + 1}, colunas "
                            f"'{cab[i_tj]}' / '{cab[i_login]}' / '{cab[i_senha]}'",
                        )
        finally:
            wb.close()
    if melhor is None:
        raise RuntimeError(
            f"nenhuma aba com colunas de tribunal/login/senha (abas: {abas})"
        )
    return melhor


def coletar_credenciais_da_planilha(
    log: Callable[[str], None] = print,
) -> list[dict[str, str]]:
    """Lê a planilha de senhas do escritório. Somente leitura.

    Estoura ``RuntimeError`` em qualquer falha (sem acesso, sem internet, sem as
    colunas, tabela curta demais, coluna de senha vazia) — mesmo princípio de
    `coletar_processos_das_planilhas`: falha de acesso NÃO pode ser confundida
    com "as senhas já estão em dia". Quem chama traduz com
    `explicar_falha_de_acesso`.
    """
    s, robot = _sessao()
    log(f"🔑 conta de serviço: {robot} (a planilha de senhas precisa estar "
        "compartilhada como Leitor com ele)")
    itens, como = ler_senhas_da_planilha(s, SHEET_SENHAS, log=log)
    log(f"🔐 planilha de senhas: {len(itens)} linha(s) — {como}")
    return itens


ROBOT_PADRAO = "att-192@projetofelipe-503412.iam.gserviceaccount.com"


def explicar_falha_de_acesso(msg: str, robot: str = "") -> str:
    """Traduz o erro cru do Google/rede em algo que o escritório resolve.

    Mora aqui, e não na GUI, para ficar coberta por teste. Como a planilha é a
    ÚNICA fonte de senhas, esta mensagem é a única saída do usuário quando algo
    falha — "404 File not found" não diz a ninguém que falta compartilhar.
    """
    m = str(msg or "")
    baixo = m.lower()
    conta = robot or ROBOT_PADRAO
    compartilhar = (
        "A planilha de senhas não foi encontrada ou não está compartilhada com "
        f"o robô.\n\nAbra a planilha no Google → Compartilhar → adicione\n{conta}\n"
        "como Leitor, e tente de novo."
    )
    if "conta de serviço não encontrada" in baixo:
        return ("Este programa foi montado sem a credencial do Google "
                "(token_google_api.json). Avise quem gerou esta versão.")
    if "nenhuma aba com colunas" in baixo:
        return ("Encontrei a planilha, mas não achei as colunas de Tribunal, "
                "Login e Senha.\n\nA planilha precisa de uma linha de cabeçalho "
                "com esses três nomes (a antiga usava tj, Login, Senha) e pelo "
                f"menos {MIN_LINHAS_SENHAS} linhas preenchidas.")
    if any(t in baixo for t in ("accessnotconfigured", "has not been used")):
        return ("As APIs Google Drive/Sheets não estão habilitadas no projeto "
                f"da conta {conta}.")
    if any(t in baixo for t in ("404", "not found", "notfound")):
        return compartilhar
    if any(t in baixo for t in ("403", "forbidden", "permission")):
        return compartilhar
    if any(t in baixo for t in ("getaddrinfo", "max retries", "connection",
                                "timed out", "timeout", "ssl")):
        return ("Sem acesso à internet (ou a rede bloqueou o Google). "
                "Conecte-se e tente de novo.")
    return m[:200]
