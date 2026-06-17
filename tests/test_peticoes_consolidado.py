# -*- coding: utf-8 -*-
"""Cobertura do consolidado de petições e caminhos de erro."""

from __future__ import annotations

from pathlib import Path

import docx
import pytest

import organizar_processos as op
from organizador.models import Contexto, Processo


def _criar_template(path: Path, header: str, body: list[str], closing: str | None = None, com_tabela: bool = False):
    doc = docx.Document()
    doc.add_paragraph(header)
    doc.add_paragraph("")
    for linha in body:
        doc.add_paragraph(linha)
    if com_tabela:
        tabela = doc.add_table(rows=1, cols=1)
        tabela.cell(0, 0).text = "Tabela (numero do processo)"
    if closing is not None:
        doc.add_paragraph(closing)
    doc.save(str(path))


def test_gerar_peticao_consolidada_preserva_tabela_e_deduplica_body(tmp_path):
    t1 = tmp_path / "a.docx"
    t2 = tmp_path / "b.docx"
    _criar_template(
        t1,
        "EXCELENTÍSSIMO(A) SENHOR(A) JUIZ(A)",
        ["Linha longa de corpo para deduplicar com mais de vinte e cinco chars.", "Corpo exclusivo A"],
        "Fecho A",
        com_tabela=True,
    )
    _criar_template(
        t2,
        "CABECALHO IGNORADO",
        ["Linha longa de corpo para deduplicar com mais de vinte e cinco chars.", "Corpo exclusivo B"],
        "Fecho B",
    )

    saida = op._gerar_peticao_consolidada(
        [t1, t2],
        tmp_path,
        {
            "numero": "4004459-04.2026.8.26.0248",
            "vara": "1ª Vara Cível",
            "comarca": "Foro Regional VI - Penha de França",
            "estado": "SP",
        },
        Contexto(genero_juiz="m"),
    )

    doc = docx.Document(str(saida))
    textos = [p.text for p in doc.paragraphs if p.text.strip()]
    assert textos[0] == (
        "EXCELENTÍSSIMO SENHOR JUIZ DE DIREITO DA 1ª VARA CÍVEL "
        "DO FORO REGIONAL VI - PENHA DE FRANÇA/SP"
    )
    assert textos.count("Linha longa de corpo para deduplicar com mais de vinte e cinco chars.") == 1
    assert "Corpo exclusivo A" in textos
    assert "Corpo exclusivo B" in textos
    assert "Fecho A" in textos
    assert len(doc.tables) == 1
    assert "4004459-04.2026.8.26.0248" in doc.tables[0].cell(0, 0).text
    assert all(r.bold for r in doc.paragraphs[0].runs if r.text)


def test_gerar_peticao_consolidada_preserva_bloco_de_fechamento_inteiro(tmp_path):
    # O fechamento real tem vários parágrafos (fórmula + local/data + assinatura
    # + OAB). O consolidado deve emiti-los todos, não só o primeiro.
    t1 = tmp_path / "a.docx"
    doc = docx.Document()
    doc.add_paragraph("EXCELENTÍSSIMO(A) SENHOR(A) JUIZ(A)")
    doc.add_paragraph("")
    doc.add_paragraph("Corpo comprido o suficiente para não ser descartado pela dedup.")
    doc.add_paragraph("Nesses termos, pede deferimento.")
    doc.add_paragraph("De Belo Horizonte/MG para São Paulo/SP, data.")
    doc.add_paragraph("Pp. Fulano de Tal")
    doc.add_paragraph("OAB/MG 104.784")
    doc.save(str(t1))

    t2 = tmp_path / "b.docx"
    _criar_template(t2, "CABECALHO", ["Outro corpo comprido o suficiente aqui dentro."])

    saida = op._gerar_peticao_consolidada([t1, t2], tmp_path, {}, Contexto(genero_juiz="m"))
    textos = [p.text for p in docx.Document(str(saida)).paragraphs]
    for esperado in (
        "Nesses termos, pede deferimento.",
        "De Belo Horizonte/MG para São Paulo/SP, data.",
        "Pp. Fulano de Tal",
        "OAB/MG 104.784",
    ):
        assert esperado in textos, f"fechamento perdeu: {esperado!r}"


def test_gerar_peticao_consolidada_injeta_closing_default(tmp_path):
    t1 = tmp_path / "a.docx"
    _criar_template(t1, "EXCELENTÍSSIMO(A) SENHOR(A) JUIZ(A)", ["Corpo sem fechamento suficiente e comprido."])
    saida = op._gerar_peticao_consolidada([t1], tmp_path, {}, Contexto(genero_juiz="f"))
    doc = docx.Document(str(saida))
    assert any("Termos em que pede deferimento." == p.text for p in doc.paragraphs)


def test_gerar_peticao_consolidada_sem_templates_falha(tmp_path):
    with pytest.raises(ValueError, match="Nenhum template selecionado"):
        op._gerar_peticao_consolidada([], tmp_path, {}, None)


def test_gerar_peticao_unica_formato_invalido_falha(tmp_path):
    template = tmp_path / "modelo.txt"
    template.write_text("x", encoding="utf-8")
    with pytest.raises(ValueError, match="Formato de template não suportado"):
        op._gerar_peticao_unica(template, tmp_path, {})


def test_gerar_peticao_unica_doc_legado_falha(tmp_path):
    template_doc = tmp_path / "modelo.doc"
    template_doc.write_text("binario", encoding="utf-8")
    with pytest.raises(ValueError, match="Use apenas .docx"):
        op._gerar_peticao_unica(template_doc, tmp_path, {"numero": "123"})


def test_gerar_peticoes_modo_documento_unico_registra_erro(monkeypatch, tmp_path):
    proc = Processo(numero="123")
    logs = []
    monkeypatch.setattr(op.peticoes, "_gerar_peticao_consolidada", lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("quebrou")))
    op.gerar_peticoes(proc, [tmp_path / "a.docx", tmp_path / "b.docx"], tmp_path, {}, log=logs.append, documento_unico=True)
    assert any("ERRO ao consolidar" in linha for linha in logs)


def test_gerar_peticoes_modo_documento_unico_loga_sucesso(monkeypatch, tmp_path):
    proc = Processo(numero="123")
    logs = []
    monkeypatch.setattr(op.peticoes, "_gerar_peticao_consolidada", lambda *args, **kwargs: tmp_path / "saida.docx")
    op.gerar_peticoes(proc, [tmp_path / "a.docx", tmp_path / "b.docx"], tmp_path, {}, log=logs.append, documento_unico=True)
    assert any("OK: documento consolidado -> saida.docx" in linha for linha in logs)
