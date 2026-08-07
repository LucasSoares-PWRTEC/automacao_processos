# -*- coding: utf-8 -*-
"""Leitura de credenciais locais do projeto.

Formato principal: ``config/credenciais.json`` com uma lista em ``credenciais``:

{
  "credenciais": [
    {"tj": "TJSP - Esaj", "login": "...", "senha": "..."}
  ]
}

Fallback temporário: ``config/senhas.xlsx`` no layout legado.
"""

from __future__ import annotations

import json
import re
import shutil
from functools import lru_cache
from pathlib import Path

from unidecode import unidecode

from ...config import CONFIG_DIR, CREDENCIAIS_EXAMPLE_PATH, CREDENCIAIS_PATH, LEGACY_SENHAS_XLSX_PATH

SENHAS_XLSX_PATH = CONFIG_DIR / "senhas.xlsx"


def _norm(s: str) -> str:
    """minúsculo, sem acento, espaços/hífens colapsados."""
    s = unidecode(str(s or "")).lower()
    s = re.sub(r"[\s\-–—]+", " ", s)
    return s.strip()


def _to_str(valor) -> str:
    """Normaliza o valor de login/senha para string utilizável."""
    if valor is None:
        return ""
    s = str(valor).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s


def _bootstrap_credenciais_locais() -> None:
    """Importa credenciais legadas para `config/` quando possível."""
    CONFIG_DIR.mkdir(parents=True, exist_ok=True)
    if CREDENCIAIS_PATH.exists() or SENHAS_XLSX_PATH.exists():
        return
    if LEGACY_SENHAS_XLSX_PATH.exists():
        shutil.copy2(LEGACY_SENHAS_XLSX_PATH, SENHAS_XLSX_PATH)


@lru_cache(maxsize=1)
def _carregar_credenciais() -> list[dict[str, str]]:
    _bootstrap_credenciais_locais()
    if CREDENCIAIS_PATH.exists():
        payload = json.loads(CREDENCIAIS_PATH.read_text(encoding="utf-8"))
        itens = payload.get("credenciais", payload)
        if not isinstance(itens, list):
            raise ValueError(
                f"Formato inválido em {CREDENCIAIS_PATH}. "
                "Use uma lista de credenciais ou a chave 'credenciais'."
            )
        saida: list[dict[str, str]] = []
        for item in itens:
            if not isinstance(item, dict):
                continue
            tj = _to_str(item.get("tj"))
            if not tj:
                continue
            saida.append(
                {
                    "tj": tj,
                    "login": _to_str(item.get("login")),
                    "senha": _to_str(item.get("senha")),
                    "_tj_norm": _norm(tj),
                }
            )
        return saida

    if SENHAS_XLSX_PATH.exists():
        # openpyxl em vez de pandas: é o mesmo motor que o pandas usaria para .xlsx,
        # sem arrastar numpy/pyarrow para o bundle (ver OrganizadorProcessos.spec).
        from openpyxl import load_workbook

        wb = load_workbook(SENHAS_XLSX_PATH, read_only=True, data_only=True)
        try:
            linhas = wb[wb.sheetnames[0]].iter_rows(values_only=True)
            cabecalho = [str(c).strip() if c is not None else "" for c in next(linhas, ())]
            idx = {nome: i for i, nome in enumerate(cabecalho)}
            faltando = [c for c in ("tj", "Login", "Senha") if c not in idx]
            if faltando:
                raise KeyError(
                    f"{SENHAS_XLSX_PATH}: coluna(s) ausente(s) {faltando} "
                    f"(cabeçalho lido: {cabecalho})"
                )

            def celula(linha: tuple, coluna: str) -> str:
                i = idx[coluna]
                return _to_str(linha[i]) if i < len(linha) else ""

            saida = []
            for linha in linhas:
                tj = celula(linha, "tj")
                if not tj:  # equivale ao dropna(subset=["tj"]) do pandas
                    continue
                saida.append({
                    "tj": tj,
                    "login": celula(linha, "Login"),
                    "senha": celula(linha, "Senha"),
                    "_tj_norm": _norm(tj),
                })
            return saida
        finally:
            wb.close()

    raise FileNotFoundError(
        "Credenciais não encontradas. Configure "
        f"{CREDENCIAIS_PATH} (use {CREDENCIAIS_EXAMPLE_PATH} como modelo) "
        f"ou {SENHAS_XLSX_PATH}."
    )


def obter_credenciais(tj_chave: str) -> tuple[str, str]:
    """Retorna ``(login, senha)`` para a chave de tribunal informada."""
    itens = _carregar_credenciais()
    alvo = _norm(tj_chave)

    exatos = [item for item in itens if item["_tj_norm"] == alvo]
    if len(exatos) == 1:
        linha = exatos[0]
    else:
        contidos = [item for item in itens if alvo in item["_tj_norm"]]
        if len(contidos) == 1:
            linha = contidos[0]
        elif len(contidos) == 0:
            raise KeyError(f"Nenhuma credencial casa com tj='{tj_chave}'.")
        else:
            achados = ", ".join(item["tj"] for item in contidos)
            raise KeyError(
                f"Ambiguidade para tj='{tj_chave}'. Candidatos: {achados}. "
                "Use uma chave mais específica."
            )

    return linha["login"], linha["senha"]


def listar_credenciais() -> list[dict[str, str]]:
    """Lista as credenciais (tj, login, senha) para edição na GUI.

    Diferente de ``_carregar_credenciais``, não levanta se o arquivo ainda não
    existe: nesse caso devolve uma lista vazia (a tela mostra um aviso).
    """
    try:
        itens = _carregar_credenciais()
    except FileNotFoundError:
        return []
    return [
        {"tj": item["tj"], "login": item["login"], "senha": item["senha"]}
        for item in itens
    ]


def salvar_credenciais(itens: list[dict[str, str]]) -> None:
    """Persiste a lista de credenciais em ``config/credenciais.json``.

    Grava no formato canônico (``{"credenciais": [...]}``) e invalida o cache de
    ``_carregar_credenciais`` para que a próxima busca já use os novos valores.
    """
    CONFIG_DIR.mkdir(parents=True, exist_ok=True)
    payload = {
        "credenciais": [
            {
                "tj": _to_str(item.get("tj")),
                "login": _to_str(item.get("login")),
                "senha": _to_str(item.get("senha")),
            }
            for item in itens
            if _to_str(item.get("tj"))
        ]
    }
    CREDENCIAIS_PATH.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    _carregar_credenciais.cache_clear()


def cpf_formatado(login: str) -> str:
    """Formata CPF quando o login vier apenas com dígitos."""
    d = re.sub(r"\D", "", str(login or "")).zfill(11)
    if len(d) != 11:
        return login
    return f"{d[0:3]}.{d[3:6]}.{d[6:9]}-{d[9:11]}"
